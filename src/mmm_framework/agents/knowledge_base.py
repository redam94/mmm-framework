"""Knowledge-base ingestion + retrieval for the agent.

Pipeline: extract text from an uploaded file → chunk it → embed each chunk →
store chunks + float32 embedding blobs in ``kb_chunks`` (scoped to a project).
Retrieval embeds the query and does a brute-force cosine search over the
project's chunks (numpy), avoiding any vector-store dependency.

Optional extractors (pdf/docx) degrade gracefully: a missing dependency marks
the document ``error`` with a clear message rather than crashing ingest.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

import numpy as np

from mmm_framework.agents.embeddings import embed_documents, embed_query
from mmm_framework.api import sessions as sessions_store

# Map file extension -> document kind.
_KIND_BY_EXT = {
    ".txt": "text",
    ".md": "markdown",
    ".markdown": "markdown",
    ".csv": "csv",
    ".tsv": "csv",
    ".pdf": "pdf",
    ".docx": "docx",
    ".xlsx": "xlsx",
    ".xls": "xlsx",
    ".json": "text",
    ".log": "text",
    ".py": "text",
}


def kind_for(path: str | Path) -> str:
    return _KIND_BY_EXT.get(Path(path).suffix.lower(), "text")


# ── Text extraction ─────────────────────────────────────────────────────────


def extract_text(path: str | Path, kind: str | None = None) -> str:
    """Extract plain text from a supported file. Raises with a clear message
    when an optional dependency for pdf/docx is missing."""
    p = Path(path)
    kind = kind or kind_for(p)

    if kind in ("text", "markdown", "csv"):
        return p.read_text(encoding="utf-8", errors="replace")

    if kind == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("PDF support requires 'pypdf' (uv add pypdf).") from exc
        reader = PdfReader(str(p))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)

    if kind == "docx":
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError(
                "DOCX support requires 'python-docx' (uv add python-docx)."
            ) from exc
        document = docx.Document(str(p))
        return "\n".join(par.text for par in document.paragraphs)

    if kind == "xlsx":
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("XLSX support requires 'openpyxl'.") from exc
        wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
        parts: list[str] = []
        for ws in wb.worksheets:
            parts.append(f"# Sheet: {ws.title}")
            for row in ws.iter_rows(values_only=True):
                cells = [str(v) for v in row if v is not None]
                if cells:
                    parts.append("\t".join(cells))
        return "\n".join(parts)

    # Fallback: try utf-8 text
    return p.read_text(encoding="utf-8", errors="replace")


# ── Chunking ────────────────────────────────────────────────────────────────


def chunk_text(text: str, size: int = 1200, overlap: int = 200) -> list[str]:
    """Paragraph-aware character chunking with overlap."""
    text = (text or "").strip()
    if not text:
        return []
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    chunks: list[str] = []
    buf = ""
    for para in paragraphs:
        if len(buf) + len(para) + 2 <= size:
            buf = f"{buf}\n\n{para}" if buf else para
        else:
            if buf:
                chunks.append(buf)
            if len(para) <= size:
                buf = para
            else:
                # paragraph itself exceeds size — hard-split with overlap
                start = 0
                while start < len(para):
                    chunks.append(para[start : start + size])
                    start += size - overlap
                buf = ""
    if buf:
        chunks.append(buf)

    # add overlap between adjacent chunks for retrieval recall
    if overlap > 0 and len(chunks) > 1:
        overlapped = [chunks[0]]
        for i in range(1, len(chunks)):
            tail = chunks[i - 1][-overlap:]
            overlapped.append((tail + "\n" + chunks[i]) if tail else chunks[i])
        return overlapped
    return chunks


# ── Embedding (blob) helpers ────────────────────────────────────────────────


def _to_blob(vec: list[float]) -> tuple[bytes, int]:
    arr = np.asarray(vec, dtype=np.float32)
    return arr.tobytes(), int(arr.shape[0])


def _from_blob(blob: bytes, dim: int) -> np.ndarray:
    return np.frombuffer(blob, dtype=np.float32, count=dim)


# ── Ingest + search ─────────────────────────────────────────────────────────


def ingest_document(doc_id: str) -> dict[str, Any]:
    """Extract → chunk → embed → store. Updates the document status in place.

    ``doc_id`` must already exist (created with status ``pending``). Designed to
    run in a threadpool from the API endpoint.
    """
    doc = sessions_store.get_kb_document(doc_id)
    if doc is None:
        raise ValueError(f"kb document not found: {doc_id}")
    try:
        text = extract_text(doc["path"], doc["kind"])
        chunks = chunk_text(text)
        if not chunks:
            sessions_store.set_kb_document_status(doc_id, "ready", n_chunks=0)
            return sessions_store.get_kb_document(doc_id)  # type: ignore[return-value]

        vectors = embed_documents(chunks)
        rows: list[tuple[int, str, bytes, int]] = []
        for idx, (chunk, vec) in enumerate(zip(chunks, vectors)):
            blob, dim = _to_blob(vec)
            rows.append((idx, chunk, blob, dim))
        sessions_store.add_kb_chunks(doc["id"], doc["project_id"], rows)
        sessions_store.set_kb_document_status(doc_id, "ready", n_chunks=len(rows))
    except Exception as exc:  # noqa: BLE001 — record the failure on the doc
        sessions_store.set_kb_document_status(doc_id, "error", error=str(exc))
    return sessions_store.get_kb_document(doc_id)  # type: ignore[return-value]


def search(project_id: str, query: str, top_k: int = 6) -> list[dict[str, Any]]:
    """Brute-force cosine retrieval over a project's KB chunks."""
    chunks = sessions_store.iter_kb_chunks(project_id)
    if not chunks:
        return []

    qvec = np.asarray(embed_query(query), dtype=np.float32)
    qnorm = np.linalg.norm(qvec) or 1.0

    # name lookup for source attribution
    names = {d["id"]: d["name"] for d in sessions_store.list_kb_documents(project_id)}

    scored: list[tuple[float, dict[str, Any]]] = []
    for ch in chunks:
        vec = _from_blob(ch["embedding"], ch["dim"])
        if vec.shape != qvec.shape:
            continue
        denom = (np.linalg.norm(vec) or 1.0) * qnorm
        score = float(np.dot(vec, qvec) / denom)
        scored.append((score, ch))

    scored.sort(key=lambda t: t[0], reverse=True)
    results = []
    for score, ch in scored[: max(1, top_k)]:
        results.append(
            {
                "document": names.get(ch["document_id"], "?"),
                "document_id": ch["document_id"],
                "chunk_index": ch["chunk_index"],
                "text": ch["text"],
                "score": round(score, 4),
            }
        )
    return results
